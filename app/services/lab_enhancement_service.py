"""SQL Lab 增强服务：保存查询、异步导出、AI 反馈、分析会话、JOIN 推荐"""
import csv
import json
import logging
import os
import re
import time
from datetime import datetime, timedelta
from typing import Any, Dict, List, Optional

import aiomysql

from app.core.database import get_db_connection
from app.services.datasource_service import DataSourceService
from app.services.data_adapter.factory import get_adapter
from app.services.masking_service import MaskingService

logger = logging.getLogger(__name__)

EXPORT_DIR = os.environ.get("LAB_EXPORT_DIR", "/tmp/lab_exports")
EXPORT_MAX_ROWS = int(os.environ.get("LAB_EXPORT_MAX_ROWS", "50000"))


class LabEnhancementService:
    # ---------- Saved Queries ----------

    @staticmethod
    async def list_saved_queries(user_id: int, source_id: Optional[int] = None) -> List[Dict[str, Any]]:
        async with get_db_connection() as conn:
            async with conn.cursor(aiomysql.DictCursor) as cursor:
                sql = """
                    SELECT id, user_id, name, sql_text, source_id, lab_mode, test_params, tags, is_shared,
                           created_at, updated_at
                    FROM lab_saved_queries
                    WHERE user_id = %s OR is_shared = 1
                """
                params: list = [user_id]
                if source_id:
                    sql += " AND source_id = %s"
                    params.append(source_id)
                sql += " ORDER BY updated_at DESC LIMIT 200"
                await cursor.execute(sql, tuple(params))
                rows = await cursor.fetchall()
                for r in rows:
                    for k in ("test_params", "tags"):
                        if r.get(k) and isinstance(r[k], str):
                            try:
                                r[k] = json.loads(r[k])
                            except Exception:
                                pass
                return rows

    @staticmethod
    async def create_saved_query(user_id: int, data: Dict[str, Any]) -> int:
        async with get_db_connection() as conn:
            async with conn.cursor() as cursor:
                await cursor.execute(
                    """
                    INSERT INTO lab_saved_queries
                    (user_id, name, sql_text, source_id, lab_mode, test_params, tags, is_shared)
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
                    """,
                    (
                        user_id,
                        data["name"],
                        data["sql"],
                        data["source_id"],
                        data.get("lab_mode", "analyst"),
                        json.dumps(data.get("test_params") or {}),
                        json.dumps(data.get("tags") or []),
                        1 if data.get("is_shared") else 0,
                    ),
                )
                return cursor.lastrowid

    @staticmethod
    async def update_saved_query(user_id: int, query_id: int, data: Dict[str, Any]) -> bool:
        async with get_db_connection() as conn:
            async with conn.cursor() as cursor:
                await cursor.execute(
                    "SELECT user_id FROM lab_saved_queries WHERE id = %s", (query_id,)
                )
                row = await cursor.fetchone()
                if not row or row[0] != user_id:
                    return False
                await cursor.execute(
                    """
                    UPDATE lab_saved_queries
                    SET name=%s, sql_text=%s, source_id=%s, lab_mode=%s,
                        test_params=%s, tags=%s, is_shared=%s
                    WHERE id=%s
                    """,
                    (
                        data.get("name"),
                        data.get("sql"),
                        data.get("source_id"),
                        data.get("lab_mode", "analyst"),
                        json.dumps(data.get("test_params") or {}),
                        json.dumps(data.get("tags") or []),
                        1 if data.get("is_shared") else 0,
                        query_id,
                    ),
                )
                return True

    @staticmethod
    async def delete_saved_query(user_id: int, query_id: int) -> bool:
        async with get_db_connection() as conn:
            async with conn.cursor() as cursor:
                await cursor.execute(
                    "DELETE FROM lab_saved_queries WHERE id=%s AND user_id=%s",
                    (query_id, user_id),
                )
                return cursor.rowcount > 0

    # ---------- Table Favorites ----------

    @staticmethod
    async def list_table_favorites(user_id: int, source_id: int) -> List[Dict[str, Any]]:
        async with get_db_connection() as conn:
            async with conn.cursor(aiomysql.DictCursor) as cursor:
                await cursor.execute(
                    """
                    SELECT id, source_id, table_name, is_pinned, note, created_at, updated_at
                    FROM lab_table_favorites
                    WHERE user_id=%s AND source_id=%s
                    ORDER BY is_pinned DESC, updated_at DESC
                    """,
                    (user_id, source_id),
                )
                rows = await cursor.fetchall()
                for r in rows:
                    r["is_pinned"] = bool(r.get("is_pinned"))
                    for k in ("created_at", "updated_at"):
                        if r.get(k) and hasattr(r[k], "strftime"):
                            r[k] = r[k].strftime("%Y-%m-%d %H:%M:%S")
                return rows

    @staticmethod
    async def upsert_table_favorite(user_id: int, data: Dict[str, Any]) -> int:
        async with get_db_connection() as conn:
            async with conn.cursor() as cursor:
                await cursor.execute(
                    """
                    INSERT INTO lab_table_favorites
                    (user_id, source_id, table_name, is_pinned, note)
                    VALUES (%s, %s, %s, %s, %s)
                    ON DUPLICATE KEY UPDATE
                        is_pinned=VALUES(is_pinned),
                        note=VALUES(note),
                        updated_at=CURRENT_TIMESTAMP
                    """,
                    (
                        user_id,
                        data["source_id"],
                        data["table_name"],
                        1 if data.get("is_pinned") else 0,
                        (data.get("note") or "").strip() or None,
                    ),
                )
                if cursor.lastrowid:
                    return cursor.lastrowid
                await cursor.execute(
                    """
                    SELECT id FROM lab_table_favorites
                    WHERE user_id=%s AND source_id=%s AND table_name=%s
                    """,
                    (user_id, data["source_id"], data["table_name"]),
                )
                row = await cursor.fetchone()
                return row[0] if row else 0

    @staticmethod
    async def delete_table_favorite(user_id: int, source_id: int, table_name: str) -> bool:
        async with get_db_connection() as conn:
            async with conn.cursor() as cursor:
                await cursor.execute(
                    """
                    DELETE FROM lab_table_favorites
                    WHERE user_id=%s AND source_id=%s AND table_name=%s
                    """,
                    (user_id, source_id, table_name),
                )
                return cursor.rowcount > 0

    # ---------- Table Explorer (keyword search) ----------

    @staticmethod
    async def aggregate_table_tags(source_id: int) -> List[Dict[str, Any]]:
        """聚合摸排标签，供探索器左侧导航"""
        async with get_db_connection() as conn:
            async with conn.cursor(aiomysql.DictCursor) as cursor:
                await cursor.execute(
                    """
                    SELECT jt.tag AS name, COUNT(DISTINCT p.id) AS count
                    FROM db_table_profiles p
                    JOIN JSON_TABLE(
                        COALESCE(p.ai_tags, JSON_ARRAY()),
                        '$[*]' COLUMNS (tag VARCHAR(100) PATH '$')
                    ) jt
                    WHERE p.connection_id = %s
                      AND p.is_ignored = 0
                      AND jt.tag IS NOT NULL
                      AND TRIM(jt.tag) != ''
                    GROUP BY jt.tag
                    ORDER BY count DESC, jt.tag ASC
                    LIMIT 200
                    """,
                    (source_id,),
                )
                rows = await cursor.fetchall()
                return [{"name": r["name"], "count": int(r["count"] or 0)} for r in rows]

    @staticmethod
    async def search_tables(
        user_id: int,
        source_id: int,
        q: Optional[str] = None,
        tag: Optional[str] = None,
        scope: str = "all",
        recent_tables: Optional[List[str]] = None,
        page: int = 1,
        page_size: int = 40,
        include_ignored: bool = False,
    ) -> Dict[str, Any]:
        """关键词搜索表画像（分页），支持收藏/最近/标签筛选"""
        page_size = min(max(int(page_size), 1), 100)
        page = max(int(page), 1)
        offset = (page - 1) * page_size

        ignored_clause = "" if include_ignored else " AND p.is_ignored = 0"
        join_clause = f"""
            FROM db_table_profiles p
            LEFT JOIN lab_table_favorites f
              ON f.user_id = %s AND f.source_id = %s
             AND f.table_name COLLATE utf8mb4_unicode_ci = p.table_name
            WHERE p.connection_id = %s{ignored_clause}
        """
        base_params: List[Any] = [user_id, source_id, source_id]
        conditions: List[str] = []
        cond_params: List[Any] = []

        scope = (scope or "all").lower()
        if scope == "profiled":
            conditions.append("p.status = %s")
            cond_params.append(2)
        elif scope == "favorites":
            conditions.append("f.id IS NOT NULL")
        elif scope == "recent":
            names = [t.strip() for t in (recent_tables or []) if t and t.strip()]
            if not names:
                return {"total": 0, "page": page, "page_size": page_size, "items": []}
            placeholders = ", ".join(["%s"] * len(names))
            conditions.append(f"p.table_name IN ({placeholders})")
            cond_params.extend(names)

        if tag and tag.strip():
            conditions.append("JSON_CONTAINS(p.ai_tags, JSON_QUOTE(%s), '$')")
            cond_params.append(tag.strip())

        q_clean = (q or "").strip()
        if q_clean:
            like = f"%{q_clean}%"
            conditions.append(
                """(
                    p.table_name LIKE %s OR p.ai_term LIKE %s OR p.ai_description LIKE %s
                    OR f.note LIKE %s
                    OR JSON_SEARCH(p.ai_tags, 'one', %s, NULL, '$[*]') IS NOT NULL
                )"""
            )
            cond_params.extend([like, like, like, like, f"%{q_clean}%"])

        where_extra = (" AND " + " AND ".join(conditions)) if conditions else ""

        if q_clean:
            exact = q_clean
            prefix = f"{q_clean}%"
            like = f"%{q_clean}%"
            order_sql = """
                (CASE
                  WHEN p.table_name = %s THEN 0
                  WHEN p.table_name LIKE %s THEN 1
                  WHEN p.ai_term LIKE %s THEN 2
                  WHEN JSON_SEARCH(p.ai_tags, 'one', %s, NULL, '$[*]') IS NOT NULL THEN 3
                  WHEN p.ai_description LIKE %s THEN 4
                  WHEN f.note LIKE %s THEN 5
                  ELSE 6
                END) ASC,
                p.confidence_score DESC,
                p.table_name ASC
            """
            order_params = [exact, prefix, like, like, like, like]
        else:
            order_sql = "p.confidence_score DESC, p.table_name ASC"
            order_params = []

        count_sql = f"SELECT COUNT(*) AS cnt {join_clause}{where_extra}"
        list_sql = f"""
            SELECT p.table_name, p.table_type, p.ai_term, p.ai_description, p.ai_tags,
                   p.status, p.confidence_score, p.is_temporary, p.is_ignored, p.confidence_reason,
                   f.is_pinned, f.note AS favorite_note,
                   (CASE WHEN f.id IS NOT NULL THEN 1 ELSE 0 END) AS is_favorite
            {join_clause}{where_extra}
            ORDER BY {order_sql}
            LIMIT %s OFFSET %s
        """

        async with get_db_connection() as conn:
            async with conn.cursor(aiomysql.DictCursor) as cursor:
                count_params = base_params + cond_params
                await cursor.execute(count_sql, tuple(count_params))
                total_row = await cursor.fetchone()
                total = int(total_row["cnt"] if total_row else 0)

                list_params = base_params + cond_params + order_params + [page_size, offset]
                await cursor.execute(list_sql, tuple(list_params))
                rows = await cursor.fetchall()

        items: List[Dict[str, Any]] = []
        for row in rows:
            ai_tags = row.get("ai_tags")
            if isinstance(ai_tags, str):
                try:
                    ai_tags = json.loads(ai_tags)
                except Exception:
                    ai_tags = []
            items.append({
                "table_name": row["table_name"],
                "table_type": row["table_type"],
                "ai_term": row.get("ai_term"),
                "ai_description": row.get("ai_description"),
                "ai_tags": ai_tags or [],
                "status": row.get("status"),
                "confidence_score": row.get("confidence_score"),
                "is_temporary": bool(row.get("is_temporary")),
                "is_ignored": bool(row.get("is_ignored")),
                "confidence_reason": row.get("confidence_reason"),
                "is_favorite": bool(row.get("is_favorite")),
                "is_pinned": bool(row.get("is_pinned")),
                "favorite_note": row.get("favorite_note"),
            })

        return {
            "total": total,
            "page": page,
            "page_size": page_size,
            "items": items,
        }

    # ---------- Export Jobs ----------

    @staticmethod
    async def create_export_job(user_id: int, data: Dict[str, Any]) -> int:
        async with get_db_connection() as conn:
            async with conn.cursor() as cursor:
                await cursor.execute(
                    """
                    INSERT INTO lab_export_jobs
                    (user_id, source_id, sql_text, params, format, status)
                    VALUES (%s, %s, %s, %s, %s, 0)
                    """,
                    (
                        user_id,
                        data["source_id"],
                        data["sql"],
                        json.dumps(data.get("params") or {}),
                        data.get("format", "xlsx"),
                    ),
                )
                return cursor.lastrowid

    @staticmethod
    async def list_export_jobs(user_id: int, limit: int = 30) -> List[Dict[str, Any]]:
        async with get_db_connection() as conn:
            async with conn.cursor(aiomysql.DictCursor) as cursor:
                await cursor.execute(
                    """
                    SELECT id, source_id, format, status, row_count, error_message,
                           created_at, completed_at
                    FROM lab_export_jobs
                    WHERE user_id=%s
                    ORDER BY created_at DESC
                    LIMIT %s
                    """,
                    (user_id, limit),
                )
                rows = await cursor.fetchall()
                for r in rows:
                    for k in ("created_at", "completed_at"):
                        if r.get(k) and hasattr(r[k], "strftime"):
                            r[k] = r[k].strftime("%Y-%m-%d %H:%M:%S")
                return rows

    @staticmethod
    async def get_export_job(user_id: int, job_id: int) -> Optional[Dict[str, Any]]:
        async with get_db_connection() as conn:
            async with conn.cursor(aiomysql.DictCursor) as cursor:
                await cursor.execute(
                    "SELECT * FROM lab_export_jobs WHERE id=%s AND user_id=%s",
                    (job_id, user_id),
                )
                row = await cursor.fetchone()
                if row and row.get("params") and isinstance(row["params"], str):
                    try:
                        row["params"] = json.loads(row["params"])
                    except Exception:
                        pass
                return row

    @staticmethod
    async def _set_export_status(job_id: int, status: int, **kwargs):
        fields = ["status = %s"]
        values: list = [status]
        for k, v in kwargs.items():
            fields.append(f"{k} = %s")
            values.append(v)
        values.append(job_id)
        async with get_db_connection() as conn:
            async with conn.cursor() as cursor:
                await cursor.execute(
                    f"UPDATE lab_export_jobs SET {', '.join(fields)} WHERE id=%s",
                    tuple(values),
                )

    @staticmethod
    async def run_export_job(job_id: int, user: dict):
        job = await LabEnhancementService.get_export_job(int(user["user_id"]), job_id)
        if not job:
            return
        await LabEnhancementService._set_export_status(job_id, 1)
        os.makedirs(EXPORT_DIR, exist_ok=True)
        try:
            adapter = await get_adapter(
                (await DataSourceService.get_datasource(job["source_id"])).source_name
            )
            result = await adapter.preview(
                job["sql_text"],
                limit=EXPORT_MAX_ROWS,
                params=job.get("params") or {},
                offset=0,
            )
            if await MaskingService.should_mask(user, False):
                cols = result.get("columns", [])
                col_names = [c["name"] if isinstance(c, dict) else c for c in cols]
                result["rows"] = await MaskingService.mask_list_of_lists(
                    result["rows"], col_names
                )

            col_names = [
                c["name"] if isinstance(c, dict) else str(c)
                for c in result.get("columns", [])
            ]
            rows = result.get("rows", [])
            ext = job.get("format", "xlsx")
            filename = f"lab_export_{job_id}_{int(time.time())}.{ext}"
            filepath = os.path.join(EXPORT_DIR, filename)

            if ext == "csv":
                with open(filepath, "w", newline="", encoding="utf-8-sig") as f:
                    writer = csv.writer(f)
                    writer.writerow(col_names)
                    writer.writerows(rows)
            else:
                try:
                    import openpyxl
                    wb = openpyxl.Workbook()
                    ws = wb.active
                    ws.title = "数据"
                    ws.append(col_names)
                    for row in rows:
                        ws.append(row)
                    wb.save(filepath)
                except ImportError:
                    filepath = filepath.replace(".xlsx", ".csv")
                    with open(filepath, "w", newline="", encoding="utf-8-sig") as f:
                        writer = csv.writer(f)
                        writer.writerow(col_names)
                        writer.writerows(rows)

            await LabEnhancementService._set_export_status(
                job_id,
                2,
                row_count=len(rows),
                file_path=filepath,
                completed_at=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            )
        except Exception as e:
            logger.error(f"Export job {job_id} failed: {e}", exc_info=True)
            await LabEnhancementService._set_export_status(
                job_id,
                3,
                error_message=str(e),
                completed_at=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            )

    # ---------- AI Feedback ----------

    @staticmethod
    async def save_ai_feedback(user_id: int, data: Dict[str, Any]) -> int:
        async with get_db_connection() as conn:
            async with conn.cursor() as cursor:
                await cursor.execute(
                    """
                    INSERT INTO lab_ai_feedback
                    (user_id, source_id, prompt, generated_sql, rating, execution_success)
                    VALUES (%s, %s, %s, %s, %s, %s)
                    """,
                    (
                        user_id,
                        data.get("source_id"),
                        data.get("prompt"),
                        data.get("generated_sql"),
                        data["rating"],
                        1 if data.get("execution_success") else 0,
                    ),
                )
                return cursor.lastrowid

    @staticmethod
    def _feedback_time_bounds(start_time: Optional[str], end_time: Optional[str]) -> tuple[str, str]:
        now = datetime.now()
        if start_time:
            final_start = start_time.replace("T", " ")
            if len(final_start) == 16:
                final_start += ":00"
        else:
            final_start = (now - timedelta(days=30)).strftime("%Y-%m-%d 00:00:00")
        if end_time:
            final_end = end_time.replace("T", " ")
            if len(final_end) == 16:
                final_end += ":59"
        else:
            final_end = now.strftime("%Y-%m-%d %H:%M:%S")
        return final_start, final_end

    @staticmethod
    async def list_ai_feedback(
        viewer_user_id: int,
        can_view_all: bool,
        page: int = 1,
        size: int = 20,
        rating: Optional[int] = None,
        user_name: Optional[str] = None,
        source_id: Optional[int] = None,
        start_time: Optional[str] = None,
        end_time: Optional[str] = None,
        include_stats: bool = False,
    ) -> Dict[str, Any]:
        conditions: List[str] = []
        params: List[Any] = []

        if not can_view_all:
            conditions.append("f.user_id = %s")
            params.append(viewer_user_id)
        elif user_name:
            conditions.append("u.user_name LIKE %s")
            params.append(f"%{user_name}%")

        if rating is not None:
            conditions.append("f.rating = %s")
            params.append(rating)
        if source_id is not None:
            conditions.append("f.source_id = %s")
            params.append(source_id)

        final_start, final_end = LabEnhancementService._feedback_time_bounds(start_time, end_time)
        conditions.append("f.created_at >= %s")
        params.append(final_start)
        conditions.append("f.created_at <= %s")
        params.append(final_end)

        where = "WHERE " + " AND ".join(conditions)
        base_from = """
            FROM lab_ai_feedback f
            LEFT JOIN api_users u ON u.id = f.user_id
            LEFT JOIN sys_data_source ds ON ds.id = f.source_id
        """

        async with get_db_connection() as conn:
            async with conn.cursor(aiomysql.DictCursor) as cursor:
                await cursor.execute(f"SELECT COUNT(*) AS cnt {base_from} {where}", tuple(params))
                total_row = await cursor.fetchone()
                total = int(total_row["cnt"] if total_row else 0)

                statistics = None
                if include_stats:
                    await cursor.execute(
                        f"""
                        SELECT
                            COUNT(*) AS total,
                            SUM(CASE WHEN f.rating = 2 THEN 1 ELSE 0 END) AS up_count,
                            SUM(CASE WHEN f.rating = 1 THEN 1 ELSE 0 END) AS down_count,
                            SUM(CASE WHEN f.execution_success = 1 THEN 1 ELSE 0 END) AS exec_success_count
                        {base_from} {where}
                        """,
                        tuple(params),
                    )
                    stat_row = await cursor.fetchone() or {}
                    up = int(stat_row.get("up_count") or 0)
                    down = int(stat_row.get("down_count") or 0)
                    rated = up + down
                    statistics = {
                        "total": int(stat_row.get("total") or 0),
                        "up_count": up,
                        "down_count": down,
                        "exec_success_count": int(stat_row.get("exec_success_count") or 0),
                        "satisfaction_rate": round(up / rated * 100, 1) if rated else 0,
                    }

                offset = (page - 1) * size
                list_params = list(params) + [size, offset]
                await cursor.execute(
                    f"""
                    SELECT f.id, f.user_id, u.user_name, f.source_id, ds.source_name,
                           f.prompt, f.generated_sql, f.rating, f.execution_success, f.created_at
                    {base_from} {where}
                    ORDER BY f.created_at DESC
                    LIMIT %s OFFSET %s
                    """,
                    tuple(list_params),
                )
                items = await cursor.fetchall()

        for item in items:
            if item.get("created_at"):
                item["created_at"] = item["created_at"].strftime("%Y-%m-%d %H:%M:%S")

        return {
            "total": total,
            "page": page,
            "size": size,
            "items": items,
            "statistics": statistics,
        }

    # ---------- Analysis Sessions ----------

    @staticmethod
    async def list_analysis_sessions(user_id: int) -> List[Dict[str, Any]]:
        async with get_db_connection() as conn:
            async with conn.cursor(aiomysql.DictCursor) as cursor:
                await cursor.execute(
                    """
                    SELECT id, title, sql_text, created_at
                    FROM lab_analysis_sessions
                    WHERE user_id=%s ORDER BY created_at DESC LIMIT 50
                    """,
                    (user_id,),
                )
                return await cursor.fetchall()

    @staticmethod
    async def save_analysis_session(user_id: int, data: Dict[str, Any]) -> int:
        async with get_db_connection() as conn:
            async with conn.cursor() as cursor:
                await cursor.execute(
                    """
                    INSERT INTO lab_analysis_sessions
                    (user_id, title, sql_text, columns_json, messages_json)
                    VALUES (%s, %s, %s, %s, %s)
                    """,
                    (
                        user_id,
                        data["title"],
                        data.get("sql"),
                        json.dumps(data.get("columns") or []),
                        json.dumps(data.get("messages") or []),
                    ),
                )
                return cursor.lastrowid

    @staticmethod
    async def get_analysis_session(user_id: int, session_id: int) -> Optional[Dict[str, Any]]:
        async with get_db_connection() as conn:
            async with conn.cursor(aiomysql.DictCursor) as cursor:
                await cursor.execute(
                    "SELECT * FROM lab_analysis_sessions WHERE id=%s AND user_id=%s",
                    (session_id, user_id),
                )
                row = await cursor.fetchone()
                if row:
                    for k in ("columns_json", "messages_json"):
                        if row.get(k) and isinstance(row[k], str):
                            row[k] = json.loads(row[k])
                return row

    @staticmethod
    async def delete_analysis_session(user_id: int, session_id: int) -> bool:
        async with get_db_connection() as conn:
            async with conn.cursor() as cursor:
                await cursor.execute(
                    "DELETE FROM lab_analysis_sessions WHERE id=%s AND user_id=%s",
                    (session_id, user_id),
                )
                return cursor.rowcount > 0

    # ---------- JOIN Path Recommendation ----------

    @staticmethod
    def _normalize_join_type(raw: Optional[str]) -> str:
        """元数据 join_type 可能是 LEFT 或 LEFT JOIN，统一为 LEFT/INNER/..."""
        s = (raw or "LEFT").upper().strip()
        s = re.sub(r"\s+JOIN\s*$", "", s)
        if s in ("LEFT", "RIGHT", "INNER", "FULL", "CROSS"):
            return s
        return "LEFT"

    @staticmethod
    def _build_join_snippet(join_type: str, target_table: str, condition: str) -> str:
        jtype = LabEnhancementService._normalize_join_type(join_type)
        return f"{jtype} JOIN {target_table} ON {condition}"

    @staticmethod
    async def get_join_paths(source_id: int, tables: List[str]) -> List[Dict[str, Any]]:
        if len(tables) < 2:
            return []
        ds = await DataSourceService.get_datasource(source_id)
        if not ds:
            return []

        table_set = {t.lower() for t in tables}
        async with get_db_connection() as conn:
            async with conn.cursor(aiomysql.DictCursor) as cursor:
                await cursor.execute(
                    """
                    SELECT r.id, r.join_condition, r.join_type, r.description,
                           t1.physical_name AS source_table, t2.physical_name AS target_table,
                           d.data_source
                    FROM meta_relationships r
                    JOIN meta_tables t1 ON r.source_table_id = t1.id
                    JOIN meta_tables t2 ON r.target_table_id = t2.id
                    JOIN meta_datasets d ON t1.dataset_id = d.id
                    WHERE d.data_source = %s
                    """,
                    (ds.source_name,),
                )
                rels = await cursor.fetchall()

        paths: List[Dict[str, Any]] = []
        seen: set[tuple[str, str, str]] = set()
        for rel in rels:
            src = (rel.get("source_table") or "").lower()
            tgt = (rel.get("target_table") or "").lower()
            if src in table_set and tgt in table_set:
                cond = rel.get("join_condition") or ""
                key = (src, tgt, cond.strip().lower())
                if key in seen:
                    continue
                seen.add(key)
                jtype = LabEnhancementService._normalize_join_type(rel.get("join_type"))
                snippet = LabEnhancementService._build_join_snippet(
                    rel.get("join_type"), rel["target_table"], cond
                )
                paths.append({
                    "source_table": rel["source_table"],
                    "target_table": rel["target_table"],
                    "join_type": jtype,
                    "condition": cond,
                    "description": rel.get("description") or "",
                    "snippet": snippet,
                    "confidence": 0.95,
                })

        # 同名字段启发式 JOIN（低置信度）
        if len(paths) < len(tables) - 1:
            for i, t1 in enumerate(tables):
                for t2 in tables[i + 1:]:
                    if any(p["source_table"] == t1 and p["target_table"] == t2 for p in paths):
                        continue
                    for suffix in ("id", "_id", "code"):
                        c1 = f"{t1.split('.')[-1]}_{suffix}" if suffix != "id" else "id"
                        snippet = (
                            f"LEFT JOIN {t2} ON {t1}.{c1} = {t2}.{c1}  -- 启发式关联，请核实"
                        )
                        paths.append({
                            "source_table": t1,
                            "target_table": t2,
                            "join_type": "LEFT",
                            "condition": f"{t1}.{c1} = {t2}.{c1}",
                            "description": "基于字段名启发式推荐，需人工确认",
                            "snippet": snippet,
                            "confidence": 0.4,
                        })
                        break
        return sorted(paths, key=lambda p: -p["confidence"])

    # ---------- Publish Health Check ----------

    @staticmethod
    def infer_param_schema(sql: str) -> List[Dict[str, Any]]:
        from app.utils.jinja_sql import SQL_LAB_ENV
        from jinja2 import meta

        clean = re.sub(r"--.*$", "", sql, flags=re.MULTILINE)
        clean = re.sub(r"/\*[\s\S]*?\*/", "", clean)
        params: List[Dict[str, Any]] = []
        try:
            ast = SQL_LAB_ENV.parse(clean)
            for name in sorted(meta.find_undeclared_variables(ast)):
                lower = name.lower()
                ptype = "string"
                if "time" in lower or "date" in lower:
                    ptype = "datetime"
                elif lower.endswith("_id") or lower == "id":
                    ptype = "integer"
                params.append({
                    "name": name,
                    "type": ptype,
                    "required": False,
                    "example": "",
                })
        except Exception:
            pass
        return params
